"""Сводный отчёт для акимата / министерства: Excel и PDF по данным дашборда туризма."""
import io

from ..config import APP_NAME
from ..i18n import t

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _rows(d: dict, lang: str) -> dict:
    """Одни и те же таблицы для обоих форматов."""
    c, p, tr = d["cur"], d["prev"], d["trend"]
    usd = d["usd"]

    def pct(k):
        return "—" if tr[k] is None else f"{tr[k]:+.1f}%"

    kpi = [
        (t(lang, "gov.k_gmv"), c["gmv"], round(c["gmv"] / usd) if usd else "", p["gmv"], pct("gmv")),
        (t(lang, "gov.k_local"), c["local"], round(c["local"] / usd) if usd else "", p["local"], pct("local")),
        (t(lang, "gov.k_tourists"), c["tourists"], "", p["tourists"], pct("tourists")),
        (t(lang, "gov.k_tourist_days"), c["tourist_days"], "", p["tourist_days"], pct("tourist_days")),
        (t(lang, "gov.k_avg_stay"), c["avg_stay"], "", p["avg_stay"], pct("avg_stay")),
        (t(lang, "gov.k_bookings"), c["bookings"], "", p["bookings"], ""),
        (t(lang, "gov.k_active_guides"), d["active"], "", "", ""),
        (t(lang, "gov.k_verified_guides"), d["verified"], "", "", ""),
        (t(lang, "gov.k_guide_month"), d["guide_month"], "", "", ""),
    ]
    countries = [(t(lang, "country." + x["country"]), float(x["people"])) for x in d["countries"]]
    if d["unknown"]:
        countries.append((t(lang, "gov.unknown_country"), d["unknown"]))
    sites = [(x[f"name_{lang}"], int(x["people"]), x["tours"]) for x in d["sites"]]
    months = [(x["m"].strftime("%m.%Y"), int(x["gmv"]), int(x["local"]), int(x["tourists"])) for x in d["monthly"]]
    return {"kpi": kpi, "countries": countries, "sites": sites, "months": months}


def title(d: dict, lang: str) -> str:
    return t(lang, "gov.report_title", app=APP_NAME, start=d["start"].strftime("%d.%m.%Y"), end=d["end"].strftime("%d.%m.%Y"))


def xlsx(d: dict, lang: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = _rows(d, lang)
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, head, rows):
        ws.append(head)
        for cell in ws[ws.max_row]:
            cell.font = bold
        for row in rows:
            ws.append(list(row))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))

    ws = wb.active
    ws.title = t(lang, "gov.sheet_kpi")[:31]
    ws.append([title(d, lang)]); ws["A1"].font = Font(bold=True, size=13); ws.append([])
    sheet(ws, [t(lang, "gov.col_metric"), t(lang, "gov.col_value"), "USD", t(lang, "gov.col_prev"), t(lang, "gov.col_change")], r["kpi"])
    sheet(wb.create_sheet(t(lang, "gov.sheet_countries")[:31]), [t(lang, "dash.col_country"), t(lang, "dash.col_people")], r["countries"])
    sheet(wb.create_sheet(t(lang, "gov.sheet_sites")[:31]), [t(lang, "gov.col_site"), t(lang, "dash.col_people"), t(lang, "gov.col_tours")], r["sites"])
    sheet(wb.create_sheet(t(lang, "gov.sheet_months")[:31]), [t(lang, "gov.col_month"), "GMV, ₸", t(lang, "gov.k_local") + ", ₸", t(lang, "gov.k_tourists")], r["months"])
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def _num(v) -> str:
    if isinstance(v, float) and not v.is_integer():
        return f"{v:,.1f}".replace(",", " ")
    return f"{int(v):,}".replace(",", " ") if isinstance(v, (int, float)) else str(v)


def pdf(d: dict, lang: str) -> bytes:
    from fpdf import FPDF
    r = _rows(d, lang)
    doc = FPDF()
    doc.add_font("dv", "", FONT)
    doc.add_font("dv", "B", FONT_B)
    doc.set_auto_page_break(True, 15)
    doc.add_page()
    doc.set_font("dv", "B", 14)
    doc.multi_cell(0, 8, title(d, lang), new_x="LMARGIN", new_y="NEXT")
    doc.set_font("dv", "", 9)
    doc.multi_cell(0, 5, t(lang, "gov.report_note"), new_x="LMARGIN", new_y="NEXT")
    doc.ln(3)

    def table(head, rows, widths):
        doc.set_font("dv", "B", 9)
        for h, w in zip(head, widths):
            doc.cell(w, 7, str(h), border="B")
        doc.ln()
        doc.set_font("dv", "", 9)
        for row in rows:
            for i, (v, w) in enumerate(zip(row, widths)):
                doc.cell(w, 6, _num(v) if i else str(v), align="R" if i else "L")
            doc.ln()
        doc.ln(4)

    def h(text):
        doc.set_font("dv", "B", 11); doc.cell(0, 8, text); doc.ln()

    table([t(lang, "gov.col_metric"), t(lang, "gov.col_value"), "USD", t(lang, "gov.col_prev"), t(lang, "gov.col_change")],
          r["kpi"], [70, 32, 26, 32, 26])
    h(t(lang, "gov.geo_title")); table([t(lang, "dash.col_country"), t(lang, "dash.col_people")], r["countries"], [110, 40])
    h(t(lang, "gov.sites_title")); table([t(lang, "gov.col_site"), t(lang, "dash.col_people"), t(lang, "gov.col_tours")], r["sites"], [100, 40, 40])
    h(t(lang, "gov.months_title")); table([t(lang, "gov.col_month"), "GMV, ₸", t(lang, "gov.k_local") + ", ₸", t(lang, "gov.k_tourists")], r["months"], [40, 45, 55, 40])
    return bytes(doc.output())
